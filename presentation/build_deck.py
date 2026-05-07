"""Build NOCTURNE_presentation.pptx for NBA 6925.

Five-slide deck per the course rubric (Slide 1 title; Slide 2 problem;
Slide 3 demo; Slide 4 user study; Slide 5 evals + preliminary results).
Theme matches the live UI: dark background, gold (#F2CA50) accent,
serif headers, sans-serif body. No em-dashes anywhere in user-facing copy.

Usage:
    .\\.venv\\Scripts\\python.exe presentation\\build_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ── Theme ────────────────────────────────────────────────────────────────────
BG_DARK = RGBColor(0x0F, 0x0E, 0x0E)
SURFACE = RGBColor(0x1C, 0x1B, 0x1B)
PANEL = RGBColor(0x24, 0x23, 0x22)
LINE = RGBColor(0x4D, 0x46, 0x35)
GOLD = RGBColor(0xF2, 0xCA, 0x50)
TEXT = RGBColor(0xE5, 0xE2, 0xE1)
MUTED = RGBColor(0xD0, 0xC5, 0xAF)
DIM = RGBColor(0x8A, 0x83, 0x77)
SUCCESS = RGBColor(0x16, 0xA3, 0x4A)
MISS = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)

HEADER_FONT = "Georgia"
BODY_FONT = "Calibri"

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT = PROJECT_ROOT / "NOCTURNE_presentation.pptx"


# ── Helpers ──────────────────────────────────────────────────────────────────

def set_solid_fill(shape, color: RGBColor) -> None:
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color


def set_no_line(shape) -> None:
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, color, border_color=None, border_w=None):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    set_solid_fill(rect, color)
    if border_color is None:
        set_no_line(rect)
    else:
        rect.line.color.rgb = border_color
        if border_w is not None:
            rect.line.width = border_w
    rect.shadow.inherit = False
    return rect


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size=14,
    bold=False,
    italic=False,
    color=TEXT,
    font=BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    char_spacing=None,
    line_spacing=None,
):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = valign
    if isinstance(text, str):
        text = [text]
    for i, line in enumerate(text):
        if isinstance(line, dict):
            content = line.get("text", "")
            seg_size = line.get("size", size)
            seg_bold = line.get("bold", bold)
            seg_italic = line.get("italic", italic)
            seg_color = line.get("color", color)
            seg_font = line.get("font", font)
            seg_align = line.get("align", align)
        else:
            content = line
            seg_size = size
            seg_bold = bold
            seg_italic = italic
            seg_color = color
            seg_font = font
            seg_align = align
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = seg_align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = content
        run.font.name = seg_font
        run.font.size = Pt(seg_size)
        run.font.bold = seg_bold
        run.font.italic = seg_italic
        run.font.color.rgb = seg_color
    return tb


def add_circle(slide, cx, cy, r, color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    set_solid_fill(s, color)
    set_no_line(s)
    s.shadow.inherit = False
    return s


def add_full_bg(slide, color):
    return add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), color)


# ── Build presentation (16:9 widescreen, 13.333 × 7.5 in) ────────────────────

pres = Presentation()
pres.slide_width = Inches(13.333)
pres.slide_height = Inches(7.5)
SW = pres.slide_width
SH = pres.slide_height

BLANK = pres.slide_layouts[6]


# ── Slide 1: Title ───────────────────────────────────────────────────────────

s1 = pres.slides.add_slide(BLANK)
add_full_bg(s1, BG_DARK)

# Two thin gold vertical accents on outer margins
add_rect(s1, Inches(0.55), Inches(0.7), Inches(0.04), Inches(6.1), GOLD)
add_rect(s1, Inches(12.74), Inches(0.7), Inches(0.04), Inches(6.1), GOLD)

# Eyebrow label, top-centered
add_text(
    s1, Inches(0), Inches(1.2), Inches(13.333), Inches(0.5),
    "NBA 6925  ·  Cornell SC Johnson",
    size=14, color=GOLD, font=BODY_FONT, bold=True,
    align=PP_ALIGN.CENTER, char_spacing=400,
)

# Big serif title
add_text(
    s1, Inches(0), Inches(2.0), Inches(13.333), Inches(2.2),
    "NOCTURNE",
    size=140, color=GOLD, bold=True, font=HEADER_FONT,
    align=PP_ALIGN.CENTER,
)

# Tagline
add_text(
    s1, Inches(0), Inches(4.3), Inches(13.333), Inches(0.7),
    "AI matchmaking for DJs and venues",
    size=26, color=TEXT, italic=True, font=HEADER_FONT,
    align=PP_ALIGN.CENTER,
)

# Footer block (author + date)
add_text(
    s1, Inches(0), Inches(6.2), Inches(13.333), Inches(0.4),
    "Pablo Penalba",
    size=18, color=MUTED, bold=True, font=BODY_FONT,
    align=PP_ALIGN.CENTER,
)
add_text(
    s1, Inches(0), Inches(6.6), Inches(13.333), Inches(0.4),
    "April 30, 2026",
    size=14, color=DIM, font=BODY_FONT, align=PP_ALIGN.CENTER,
)


# ── Slide 2: Problem ─────────────────────────────────────────────────────────

s2 = pres.slides.add_slide(BLANK)
add_full_bg(s2, BG_DARK)

add_text(
    s2, Inches(0.7), Inches(0.55), Inches(6), Inches(0.4),
    "THE PROBLEM",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=400,
)
add_text(
    s2, Inches(0.7), Inches(0.95), Inches(12), Inches(1.0),
    "Booking a DJ is blind.",
    size=46, color=TEXT, bold=True, font=HEADER_FONT,
)

# Three problem cards. card_h grew 2.70 -> 2.85 so the body fits with real margin.
cards = [
    (
        "Venues book without preview",
        "Promoters pick from EPKs, demo mixes, and Instagram, then hope the vibe lands. By the night of the show, course-correction is impossible.",
    ),
    (
        "DJs pitch into voids",
        "Hours spent emailing venues that never reply, or worse, venues whose crowd doesn’t fit the set. Booking signal is sparse and one-directional.",
    ),
    (
        "No preview of the night",
        "There’s no medium that lets venues, DJs, and crowds preview the actual sonic and visual character of the show in advance.",
    ),
]

card_y = Inches(2.30)
card_h = Inches(2.85)
gap = Inches(0.40)
total_w = Inches(13.333) - Inches(0.7) * 2
card_w = (total_w - gap * 2) / 3
for i, (head, body) in enumerate(cards):
    cx = Inches(0.7) + (card_w + gap) * i
    add_rect(s2, cx, card_y, card_w, card_h, PANEL)
    add_rect(s2, cx, card_y, Inches(0.5), Inches(0.08), GOLD)
    # Card number
    add_text(
        s2, cx + Inches(0.4), card_y + Inches(0.30), Inches(2), Inches(0.4),
        f"0{i + 1}",
        size=28, color=GOLD, bold=True, font=HEADER_FONT,
    )
    # Card heading: y = sy+0.80, h = 0.70 (was 0.85/0.90, overlapped body)
    add_text(
        s2, cx + Inches(0.4), card_y + Inches(0.80), card_w - Inches(0.8), Inches(0.70),
        head,
        size=18, color=TEXT, bold=True, font=HEADER_FONT,
    )
    # Card body: starts after heading (sy+1.55), h=1.20 fits ~5 lines comfortably.
    add_text(
        s2, cx + Inches(0.4), card_y + Inches(1.55), card_w - Inches(0.8), Inches(1.20),
        body,
        size=12, color=MUTED, font=BODY_FONT, line_spacing=1.25,
    )

# Bottom AI-fit strip. body h trimmed 0.95 -> 0.85 so it fits inside.
strip_y = Inches(5.45)
strip_h = Inches(1.45)
add_rect(s2, Inches(0.7), strip_y, total_w, strip_h, SURFACE)
add_rect(s2, Inches(0.7), strip_y, Inches(0.08), strip_h, GOLD)
add_text(
    s2, Inches(1.0), strip_y + Inches(0.18), Inches(8), Inches(0.4),
    "WHY AI, NOT JUST CHATGPT",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=400,
)
add_text(
    s2, Inches(1.0), strip_y + Inches(0.55), total_w - Inches(0.6), Inches(0.85),
    "Vibe is fuzzy and multi-modal: genre-tag overlap, BPM windows, room atmosphere, capacity, "
    "and bio context all need joint reasoning. ChatGPT alone has no roster, no audio analysis, "
    "and no media generation. The platform owns the data and stitches the AI stack together.",
    size=14, color=TEXT, font=BODY_FONT, line_spacing=1.3,
)


# ── Slide 3: Demo ────────────────────────────────────────────────────────────

s3 = pres.slides.add_slide(BLANK)
add_full_bg(s3, BG_DARK)

add_text(
    s3, Inches(0.7), Inches(0.55), Inches(6), Inches(0.4),
    "LIVE DEMO",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=400,
)
add_text(
    s3, Inches(0.7), Inches(0.95), Inches(12), Inches(0.95),
    "From upload to venue simulation.",
    size=42, color=TEXT, bold=True, font=HEADER_FONT,
)

# Three numbered steps. step_h 1.4 -> 1.3, gap 0.18 -> 0.15
# Total height: 3 * 1.30 + 2 * 0.15 = 4.20. Starts at 2.0 -> ends at 6.20. Bottom strip at 6.55 -> 0.35 gap.
steps = [
    (
        "Upload track + photos + style",
        "30-second audio sample (trimmed client-side), performer photos, and a one-line style description.",
    ),
    (
        "AI matches you to a venue",
        "An audio model reads the track’s style; the matcher ranks 7 venues with reasoning over genre, vibe, capacity, and bio context.",
    ),
    (
        "AI generates the night",
        "Music preview clip per matched venue plus a 10-second cinematic venue simulation conditioned on the room’s photo. Soundtrack syncs in the browser.",
    ),
]

step_x = Inches(0.7)
step_w = Inches(6.0)
step_top = Inches(2.0)
step_h = Inches(1.30)
step_gap = Inches(0.15)
for i, (head, body) in enumerate(steps):
    sy = step_top + (step_h + step_gap) * i
    add_circle(s3, step_x + Inches(0.45), sy + Inches(0.45), Inches(0.4), GOLD)
    add_text(
        s3, step_x, sy + Inches(0.05), Inches(0.9), Inches(0.8),
        f"{i + 1}",
        size=28, color=BG_DARK, bold=True, font=HEADER_FONT,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        s3, step_x + Inches(1.1), sy, step_w - Inches(1.1), Inches(0.45),
        head,
        size=18, color=TEXT, bold=True, font=HEADER_FONT,
    )
    add_text(
        s3, step_x + Inches(1.1), sy + Inches(0.48), step_w - Inches(1.1), Inches(0.80),
        body,
        size=12, color=MUTED, font=BODY_FONT, line_spacing=1.25,
    )

# Right: "watch live" stage card. Dark panel with a gold border so it visually
# pops; tells the audience the next beat is leaving PowerPoint for the browser.
ph_x = Inches(7.1)
ph_y = Inches(2.0)
ph_w = Inches(5.5)
ph_h = Inches(4.2)
add_rect(s3, ph_x, ph_y, ph_w, ph_h, SURFACE, border_color=GOLD, border_w=Pt(1.5))

# Visual hierarchy: big "WATCH LIVE" hero, smaller URL as subline.
# (Earlier draft put the URL at 36pt which wrapped to "/personal-/demo.html"
# on the 5.5" panel — flipped the hierarchy so the eye lands on the action,
# not the path.)
add_text(
    s3, ph_x, ph_y + Inches(0.85), ph_w, Inches(0.35),
    "WATCH IT LIVE",
    size=14, color=GOLD, bold=True, font=BODY_FONT, char_spacing=500,
    align=PP_ALIGN.CENTER,
)

add_text(
    s3, ph_x, ph_y + Inches(1.40), ph_w, Inches(1.05),
    "Switching to\nthe browser.",
    size=40, color=TEXT, bold=True, font=HEADER_FONT,
    align=PP_ALIGN.CENTER, line_spacing=1.05,
)

# URL subline (smaller — won't wrap on 5.5" panel)
add_text(
    s3, ph_x, ph_y + Inches(2.95), ph_w, Inches(0.40),
    "/personal-demo.html",
    size=18, color=GOLD, bold=True, font=BODY_FONT,
    align=PP_ALIGN.CENTER,
)

# Closing description
add_text(
    s3, ph_x + Inches(0.5), ph_y + Inches(3.45), ph_w - Inches(1.0), Inches(0.65),
    "Real-time venue match, AI music preview, and venue simulation.",
    size=12, color=MUTED, italic=True, font=HEADER_FONT,
    align=PP_ALIGN.CENTER, line_spacing=1.30,
)

# Bottom note - reframed for the live cut-over
note_y = Inches(6.55)
add_rect(s3, Inches(0.7), note_y, Inches(11.93), Inches(0.55), SURFACE)
add_rect(s3, Inches(0.7), note_y, Inches(0.08), Inches(0.55), GOLD)
add_text(
    s3, Inches(1.0), note_y + Inches(0.1), Inches(11.5), Inches(0.4),
    "One stack  ·  swappable models  ·  ~$0.05 per personalized run",
    size=12, color=TEXT, font=BODY_FONT, valign=MSO_ANCHOR.MIDDLE,
)


# ── Slide 4: User study ──────────────────────────────────────────────────────

s4 = pres.slides.add_slide(BLANK)
add_full_bg(s4, BG_DARK)

add_text(
    s4, Inches(0.7), Inches(0.55), Inches(6), Inches(0.4),
    "USER STUDY",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=400,
)
add_text(
    s4, Inches(0.7), Inches(0.95), Inches(12), Inches(1.0),
    "First-user walkthrough.",
    size=42, color=TEXT, bold=True, font=HEADER_FONT,
)

# Two-column layout: WHO TESTED IT (left) + WHAT'S NEXT (right).
# Quote box + 3 takeaway cards removed - it's not honest to put polished
# words in Santiago's mouth before he has actually run the demo. The
# n=1 framing + a credible plan for the broader study is the academically
# defensible move.
total_w = Inches(11.93)
col_gap = Inches(0.30)
col_w = (total_w - col_gap) / 2
col_y = Inches(2.0)
col_h = Inches(4.4)

# Left column - WHO TESTED IT
left_x = Inches(0.7)
add_rect(s4, left_x, col_y, col_w, col_h, PANEL)
add_rect(s4, left_x, col_y, Inches(0.5), Inches(0.08), GOLD)
add_text(
    s4, left_x + Inches(0.4), col_y + Inches(0.30), col_w - Inches(0.8), Inches(0.4),
    "WHO TESTED IT",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=300,
)
add_text(
    s4, left_x + Inches(0.4), col_y + Inches(0.70), col_w - Inches(0.8), Inches(0.55),
    "Santiago Vuilleumier",
    size=22, color=TEXT, bold=True, font=HEADER_FONT,
)
add_text(
    s4, left_x + Inches(0.4), col_y + Inches(1.25), col_w - Inches(0.8), Inches(0.40),
    "Miami DJ; founder & curator, Casa Whispers",
    size=13, color=MUTED, italic=True, font=HEADER_FONT,
)
add_text(
    s4, left_x + Inches(0.4), col_y + Inches(1.85), col_w - Inches(0.8), Inches(2.4),
    [
        "Protocol:",
        "  ·  Uploaded a 30-second sample of a recent set",
        "  ·  Provided photos and a free-text style description",
        "  ·  Reviewed the top-3 venue ranking and reasoning",
        "  ·  Played the AI music preview and venue simulation",
        "  ·  Talked through perceived fit and confidence after",
    ],
    size=13, color=MUTED, font=BODY_FONT, line_spacing=1.45,
)

# Right column - WHAT'S NEXT
right_x = left_x + col_w + col_gap
add_rect(s4, right_x, col_y, col_w, col_h, PANEL)
add_rect(s4, right_x, col_y, Inches(0.5), Inches(0.08), GOLD)
add_text(
    s4, right_x + Inches(0.4), col_y + Inches(0.30), col_w - Inches(0.8), Inches(0.4),
    "WHAT’S NEXT",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=300,
)
add_text(
    s4, right_x + Inches(0.4), col_y + Inches(0.70), col_w - Inches(0.8), Inches(0.55),
    "Broader study",
    size=22, color=TEXT, bold=True, font=HEADER_FONT,
)
add_text(
    s4, right_x + Inches(0.4), col_y + Inches(1.25), col_w - Inches(0.8), Inches(0.40),
    "Same protocol, wider sample, post-presentation.",
    size=13, color=MUTED, italic=True, font=HEADER_FONT,
)
add_text(
    s4, right_x + Inches(0.4), col_y + Inches(1.85), col_w - Inches(0.8), Inches(2.4),
    [
        "Plan:",
        "  ·  3 to 5 working DJs across genres",
        "  ·  2 to 3 venue managers / promoters",
        "  ·  Same protocol; blind ratings of top-3 matches",
        "  ·  Mix of audio uploads and typed-style entries",
        "  ·  Capture quotes for a follow-up writeup",
    ],
    size=13, color=MUTED, font=BODY_FONT, line_spacing=1.45,
)

# Footer note
add_text(
    s4, Inches(0.7), Inches(6.85), Inches(11.93), Inches(0.3),
    "n=1 walkthrough with Santiago Vuilleumier (Casa Whispers, Miami). Specific observations to be captured in the broader study.",
    size=10, color=DIM, italic=True, font=BODY_FONT,
)


# ── Slide 5: Evals ───────────────────────────────────────────────────────────

s5 = pres.slides.add_slide(BLANK)
add_full_bg(s5, BG_DARK)

add_text(
    s5, Inches(0.7), Inches(0.30), Inches(6), Inches(0.30),
    "EVALUATION",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=400,
)
add_text(
    s5, Inches(0.7), Inches(0.65), Inches(12.5), Inches(0.65),
    "How we know it’s working, or not.",
    size=32, color=TEXT, bold=True, font=HEADER_FONT,
)

# Layout: fixed sections leave the grid 3.00" tall so each row's body box
# (row_h - 1.0) lands at 0.40", enough for the §9.5 / §9.6 bodies to wrap to
# two lines without spilling under the closing takeaway line.
stat_y = Inches(1.45)
stat_h = Inches(1.45)
add_rect(s5, Inches(0.7), stat_y, Inches(11.93), stat_h, PANEL)
add_rect(s5, Inches(0.7), stat_y, Inches(0.08), stat_h, GOLD)
add_text(
    s5, Inches(1.0), stat_y + Inches(0.13), Inches(10), Inches(0.32),
    "§9.3  DJ-MATCH BLIND EVAL  ·  21 (VENUE, DJ) PAIRS  ·  HUMAN-RATED 1 TO 5",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=300,
)

mx = Inches(1.0)
my = stat_y + Inches(0.55)
metric_w = Inches(3.6)
metrics = [
    ("3.71 / 5.00", "Match Precision", "target ≥ 4.00", MISS),
    ("28.6%", "Rank Accuracy (top-1)", "target ≥ 70%", MISS),
    ("2 / 7", "Top-1 matched human pick", "FABRIC LONDON · THE NEST", AMBER),
]
for i, (val, label, target, color) in enumerate(metrics):
    cx = mx + (metric_w + Inches(0.2)) * i
    add_text(
        s5, cx, my, metric_w, Inches(0.50),
        val, size=32, color=color, bold=True, font=HEADER_FONT,
    )
    add_text(
        s5, cx, my + Inches(0.50), metric_w, Inches(0.22),
        label, size=12, color=TEXT, bold=True, font=BODY_FONT,
    )
    add_text(
        s5, cx, my + Inches(0.72), metric_w, Inches(0.16),
        target, size=9, color=DIM, font=BODY_FONT,
    )

# Failure mode callout (body trimmed so it fits 1-2 lines inside h=0.85).
fail_y = Inches(3.00)
fail_h = Inches(0.85)
add_rect(s5, Inches(0.7), fail_y, Inches(11.93), fail_h, SURFACE)
add_rect(s5, Inches(0.7), fail_y, Inches(0.08), fail_h, GOLD)
add_text(
    s5, Inches(1.0), fail_y + Inches(0.10), Inches(11), Inches(0.32),
    "FAILURE MODE THE EVAL CAUGHT, PRE-DEPLOYMENT",
    size=11, color=GOLD, bold=True, font=BODY_FONT, char_spacing=300,
)
add_text(
    s5, Inches(1.0), fail_y + Inches(0.42), Inches(11.5), Inches(0.40),
    "Matcher over-weights genre tags and ignores DJ bios. AELXIS VOX gets pushed to industrial "
    "venues despite a bio describing cinematic festival sets.",
    size=12, color=TEXT, font=BODY_FONT, line_spacing=1.30,
)

# Four-eval grid (2x2). Grid grew 2.80 -> 3.00 (row_h 1.30 -> 1.40, body box 0.40)
# AND body copy shortened to keep each card to 1-2 lines. Combined fix kills the
# §9.6 body overflow that was bleeding into the bottom takeaway line.
grid_y = Inches(3.95)
grid_h = Inches(3.00)
grid_gap = Inches(0.20)
grid_total_w = Inches(11.93)
col_w = (grid_total_w - grid_gap) / 2
row_h = (grid_h - grid_gap) / 2  # 1.40

eval_blocks = [
    (
        "§9.3  DJ-MATCH",
        "Live · numbers above",
        "21 blind-rated pairs. Both targets missed; failure mode caught. Iteration planned.",
        AMBER,
    ),
    (
        "§9.4  AUDIO-DRIVEN MATCH",
        "Harness ready · samples pending",
        "Paired typed-style vs typed+audio protocol on 7 venues. Sample collection in progress.",
        DIM,
    ),
    (
        "§9.5  PROMPT-STRUCTURE LINTER",
        "3 / 3 venues PASS · 15 / 15 checks each",
        "Deterministic linter for Seedance prompt structure. $0 to run, fits CI.",
        SUCCESS,
    ),
    (
        "§9.6  VIDEO OUTPUT RUBRIC",
        "1 / 3 generated · ratings pending",
        "5-dim rubric (identity, fidelity, action, camera, behavior) plus constraints check. ~$1.51 / venue.",
        AMBER,
    ),
]
for idx, (label, status, body, status_color) in enumerate(eval_blocks):
    row = idx // 2
    col = idx % 2
    cx = Inches(0.7) + (col_w + grid_gap) * col
    cy = grid_y + (row_h + grid_gap) * row
    add_rect(s5, cx, cy, col_w, row_h, PANEL)
    add_rect(s5, cx, cy, Inches(0.5), Inches(0.06), GOLD)
    add_text(
        s5, cx + Inches(0.30), cy + Inches(0.18), col_w - Inches(0.6), Inches(0.32),
        label, size=12, color=GOLD, bold=True, font=BODY_FONT, char_spacing=300,
    )
    add_text(
        s5, cx + Inches(0.30), cy + Inches(0.50), col_w - Inches(0.6), Inches(0.32),
        status, size=12, color=status_color, bold=True, font=BODY_FONT,
    )
    add_text(
        s5, cx + Inches(0.30), cy + Inches(0.85), col_w - Inches(0.6), row_h - Inches(1.0),
        body, size=11, color=MUTED, font=BODY_FONT, line_spacing=1.30,
    )

# Bottom takeaway. y=7.05 sits 0.10 below the grid (which ends at 6.95) so the
# centered text never collides with §9.5 / §9.6 bodies above it.
add_text(
    s5, Inches(0.7), Inches(7.05), Inches(11.93), Inches(0.35),
    "The eval did its job: it caught a real failure mode before any user saw it.",
    size=13, color=GOLD, bold=True, italic=True, font=HEADER_FONT,
    align=PP_ALIGN.CENTER,
)


# ── Save ─────────────────────────────────────────────────────────────────────

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
pres.save(str(OUTPUT))
print(f"Wrote {OUTPUT}")
